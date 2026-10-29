from fastapi import FastAPI, File, UploadFile, Form, Request
from fastapi.responses import HTMLResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
import pdfplumber
import shutil, re, json, io, uuid
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import spacy
from sentence_transformers import SentenceTransformer, util
import yake
from google import genai as google_genai
import os
from dotenv import load_dotenv

# Load environment variables from .env file
load_dotenv()

# =============================================
# CONFIGURATION — API KEY HIDDEN IN .env FILE
# =============================================
client = google_genai.Client(api_key=os.environ.get("GEMINI_API_KEY"))



# =============================================
# LOAD NLP MODELS
# =============================================
print("Loading NLP models...")
nlp = spacy.load("en_core_web_sm")
embedder = SentenceTransformer("all-MiniLM-L6-v2")
kw_extractor = yake.KeywordExtractor(lan="en", n=2, top=20)
print("All models loaded!")

app = FastAPI()
templates = Jinja2Templates(directory="templates")

# =============================================
# IN-MEMORY DATABASE
# =============================================
jobs_store = []
candidates_store = []
active_job_id = None

# =============================================
# HELPER FUNCTIONS
# =============================================
def extract_text(pdf_path):
    text = ""
    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            text += page.extract_text() or ""
    return text

def extract_entities(text):
    doc = nlp(text[:5000])
    entities = {"organizations": [], "dates": [], "locations": [], "persons": []}
    for ent in doc.ents:
        if ent.label_ == "ORG": entities["organizations"].append(ent.text)
        elif ent.label_ == "DATE": entities["dates"].append(ent.text)
        elif ent.label_ in ["GPE", "LOC"]: entities["locations"].append(ent.text)
        elif ent.label_ == "PERSON": entities["persons"].append(ent.text)
    return entities

def extract_keywords_yake(text):
    keywords = kw_extractor.extract_keywords(text)
    return [kw[0].lower() for kw in keywords]

def compute_semantic_similarity(resume_text, jd_text):
    resume_embedding = embedder.encode(resume_text[:1000], convert_to_tensor=True)
    jd_embedding = embedder.encode(jd_text[:1000], convert_to_tensor=True)
    similarity = util.cos_sim(resume_embedding, jd_embedding)
    return round(float(similarity[0][0]) * 100, 2)

def analyze_with_gemini(resume_text, job_description):
    prompt = f"""
You are an expert HR recruiter and ATS (Applicant Tracking System).

Analyze the resume against the job description and return ONLY a JSON object.
No markdown, no code blocks, just pure JSON.

JOB DESCRIPTION:
{job_description}

RESUME:
{resume_text[:3000]}

Return exactly this JSON:
{{
    "skill_score": <0-100>,
    "experience_score": <0-100>,
    "education_score": <0-100>,
    "experience_years": <number>,
    "education_level": "<degree>",
    "matched_skills": ["skill1", "skill2"],
    "missing_skills": ["skill1", "skill2"],
    "strengths": "<2 sentences>",
    "weaknesses": "<1 sentence>",
    "ats_score": <0-100>,
    "keyword_match_percent": <0-100>,
    "recommendation": "<Highly Recommended/Recommended/Maybe/Not Recommended>",
    "summary": "<3 sentence candidate summary>"
}}
"""
    try:
        response = client.models.generate_content(
        model="gemini-3.5-flash-lite",
        contents=prompt,
        config={
        "response_mime_type": "application/json"
    }
)
        raw = response.text.strip()
        raw = re.sub(r"```json|```", "", raw).strip()
        return json.loads(raw)
    except Exception as e:
        print(f"GEMINI ERROR: {e}")
        return {
            "skill_score": 0, "experience_score": 0, "education_score": 0,
            "experience_years": 0, "education_level": "Unknown",
            "matched_skills": [], "missing_skills": [],
            "strengths": "Could not analyze.", "weaknesses": "Could not analyze.",
            "ats_score": 0, "keyword_match_percent": 0,
            "recommendation": "Maybe", "summary": "Could not analyze resume."
        }

def extract_name(text, filename):
    doc = nlp(text[:500])
    for ent in doc.ents:
        if ent.label_ == "PERSON" and len(ent.text.split()) >= 2:
            return ent.text.title()
    lines = [l.strip() for l in text.strip().split('\n') if l.strip()]
    for line in lines[:5]:
        words = line.split()
        if 2 <= len(words) <= 4 and all(w[0].isupper() for w in words if w.isalpha()):
            return line.title()
    return filename.replace(".pdf", "").replace("_", " ").title()

def extract_email(text):
    match = re.search(r'[\w.-]+@[\w.-]+\.\w+', text)
    return match.group(0) if match else "N/A"

def extract_phone(text):
    match = re.search(r'[\+]?[(]?[0-9]{3}[)]?[-\s\.]?[0-9]{3}[-\s\.]?[0-9]{4,6}', text)
    return match.group(0) if match else "N/A"

def compute_final_score(gemini_data, semantic_score):
    skill_score = gemini_data.get("skill_score", 0)
    experience_score = gemini_data.get("experience_score", 0)
    education_score = gemini_data.get("education_score", 0)
    ats_score = gemini_data.get("ats_score", 0)
    total = round(
        skill_score      * 0.35 +
        semantic_score   * 0.25 +
        experience_score * 0.20 +
        education_score  * 0.10 +
        ats_score        * 0.10,
        2
    )
    return min(total, 100)

# =============================================
# ROUTES — DASHBOARD
# =============================================
@app.get("/", response_class=HTMLResponse)
async def dashboard(request: Request):
    total = len(candidates_store)
    selected = len([c for c in candidates_store if c["status"] == "Selected"])
    recommended = len([c for c in candidates_store if c["recommendation"] == "Highly Recommended"])
    avg_score = round(sum(c["total_score"] for c in candidates_store) / total, 1) if total else 0
    recent = sorted(candidates_store, key=lambda x: x["created_at"], reverse=True)[:5]
    return templates.TemplateResponse(request, "index.html", {
        "total": total,
        "selected": selected,
        "recommended": recommended,
        "avg_score": avg_score,
        "recent": recent,
        "jobs": jobs_store,
        "candidates": candidates_store
    })

# =============================================
# ROUTES — JOBS
# =============================================
@app.get("/jobs", response_class=HTMLResponse)
async def jobs_page(request: Request):
    return templates.TemplateResponse(request, "jobs.html", {
        "jobs": jobs_store,
        "active_job_id": active_job_id
    })

@app.post("/jobs/create", response_class=HTMLResponse)
async def create_job(
    request: Request,
    title: str = Form(...),
    company: str = Form(...),
    description: str = Form(...),
    location: str = Form(""),
    job_type: str = Form("Full-time")
):
    job = {
        "id": str(uuid.uuid4())[:8],
        "title": title,
        "company": company,
        "description": description,
        "location": location,
        "job_type": job_type,
        "created_at": datetime.now().strftime("%d %b %Y"),
        "candidates_count": 0
    }
    jobs_store.append(job)
    return templates.TemplateResponse(request, "jobs.html", {
        "jobs": jobs_store,
        "active_job_id": active_job_id,
        "success": f"Job '{title}' created successfully!"
    })

@app.post("/jobs/delete", response_class=HTMLResponse)
async def delete_job(request: Request, job_id: str = Form(...)):
    global jobs_store
    jobs_store = [j for j in jobs_store if j["id"] != job_id]
    return templates.TemplateResponse(request, "jobs.html", {
        "jobs": jobs_store,
        "active_job_id": active_job_id
    })

# =============================================
# ROUTES — SCREENING
# =============================================
@app.get("/screen", response_class=HTMLResponse)
async def screen_page(request: Request):
    return templates.TemplateResponse(request, "screen.html", {
        "jobs": jobs_store,
        "results": [],
        "job_keywords": []
    })

@app.post("/analyze", response_class=HTMLResponse)
async def analyze(
    request: Request,
    job_description: str = Form(...),
    job_id: str = Form(""),
    resumes: list[UploadFile] = File(...)
):
    jd_keywords = extract_keywords_yake(job_description)

    # Update job candidate count
    if job_id:
        for job in jobs_store:
            if job["id"] == job_id:
                job["candidates_count"] += len(resumes)

    new_candidates = []
    for resume in resumes:
        path = f"uploads/{resume.filename}"
        with open(path, "wb") as f:
            shutil.copyfileobj(resume.file, f)

        text = extract_text(path)
        entities = extract_entities(text)
        name = extract_name(text, resume.filename)
        email = extract_email(text)
        phone = extract_phone(text)
        semantic_score = compute_semantic_similarity(text, job_description)
        gemini_data = analyze_with_gemini(text, job_description)
        total_score = compute_final_score(gemini_data, semantic_score)

        existing = next((c for c in candidates_store if c["filename"] == resume.filename), None)
        if existing:
            candidates_store.remove(existing)

        candidate = {
            "id": str(uuid.uuid4())[:8],
            "name": name,
            "email": email,
            "phone": phone,
            "filename": resume.filename,
            "status": "Applied",
            "total_score": total_score,
            "semantic_score": semantic_score,
            "skill_score": gemini_data.get("skill_score", 0),
            "experience_score": gemini_data.get("experience_score", 0),
            "education_score": gemini_data.get("education_score", 0),
            "ats_score": gemini_data.get("ats_score", 0),
            "keyword_match_percent": gemini_data.get("keyword_match_percent", 0),
            "experience_years": gemini_data.get("experience_years", 0),
            "education_level": gemini_data.get("education_level", "Unknown"),
            "matched_skills": gemini_data.get("matched_skills", []),
            "missing_skills": gemini_data.get("missing_skills", []),
            "strengths": gemini_data.get("strengths", ""),
            "weaknesses": gemini_data.get("weaknesses", ""),
            "recommendation": gemini_data.get("recommendation", "Maybe"),
            "summary": gemini_data.get("summary", ""),
            "entities": entities,
            "jd_keywords": jd_keywords,
            "job_id": job_id,
            "created_at": datetime.now().strftime("%d %b %Y %H:%M"),
        }
        candidates_store.append(candidate)
        new_candidates.append(candidate)

    new_candidates.sort(key=lambda x: x["total_score"], reverse=True)
    candidates_store.sort(key=lambda x: x["total_score"], reverse=True)

    return templates.TemplateResponse(request, "screen.html", {
        "jobs": jobs_store,
        "results": new_candidates,
        "job_keywords": jd_keywords,
        "job_description": job_description
    })

# =============================================
# ROUTES — CANDIDATES
# =============================================
@app.get("/candidates", response_class=HTMLResponse)
async def candidates_page(request: Request, status: str = "", search: str = ""):
    filtered = candidates_store
    if status:
        filtered = [c for c in filtered if c["status"] == status]
    if search:
        filtered = [c for c in filtered if search.lower() in c["name"].lower()]
    return templates.TemplateResponse(request, "candidates.html", {
        "candidates": filtered,
        "status_filter": status,
        "search": search,
        "total": len(candidates_store)
    })

@app.post("/update-status", response_class=HTMLResponse)
async def update_status(
    request: Request,
    filename: str = Form(...),
    status: str = Form(...)
):
    for candidate in candidates_store:
        if candidate["filename"] == filename:
            candidate["status"] = status
            break
    filtered = candidates_store
    return templates.TemplateResponse(request, "candidates.html", {
        "candidates": filtered,
        "status_filter": "",
        "search": "",
        "total": len(candidates_store)
    })

# =============================================
# ROUTES — ANALYTICS
# =============================================
@app.get("/analytics", response_class=HTMLResponse)
async def analytics_page(request: Request):
    if not candidates_store:
        return templates.TemplateResponse(request, "analytics.html", {
            "has_data": False
        })

    scores = [c["total_score"] for c in candidates_store]
    avg_score = round(sum(scores) / len(scores), 1)
    top_score = max(scores)

    status_counts = {}
    for c in candidates_store:
        status_counts[c["status"]] = status_counts.get(c["status"], 0) + 1

    rec_counts = {}
    for c in candidates_store:
        rec_counts[c["recommendation"]] = rec_counts.get(c["recommendation"], 0) + 1

    all_matched = []
    for c in candidates_store:
        all_matched.extend(c["matched_skills"])
    skill_freq = {}
    for s in all_matched:
        skill_freq[s] = skill_freq.get(s, 0) + 1
    top_skills = sorted(skill_freq.items(), key=lambda x: x[1], reverse=True)[:10]

    return templates.TemplateResponse(request, "analytics.html", {
        "has_data": True,
        "total": len(candidates_store),
        "avg_score": avg_score,
        "top_score": top_score,
        "status_counts": status_counts,
        "rec_counts": rec_counts,
        "top_skills": top_skills,
        "candidates": candidates_store
    })

# =============================================
# ROUTES — EXPORT EXCEL
# =============================================
@app.get("/export/excel")
async def export_excel():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Candidate Rankings"

    headers = ["Rank", "Name", "Email", "Phone", "Total Score", "ATS Score",
               "Skill Score", "Semantic Score", "Experience Score", "Education Score",
               "Experience (Years)", "Education", "Recommendation", "Status",
               "Matched Skills", "Missing Skills", "Strengths", "Weaknesses"]

    header_fill = PatternFill(start_color="1a1d27", end_color="1a1d27", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for rank, c in enumerate(candidates_store, 1):
        ws.append([
            rank,
            c["name"],
            c["email"],
            c["phone"],
            c["total_score"],
            c["ats_score"],
            c["skill_score"],
            c["semantic_score"],
            c["experience_score"],
            c["education_score"],
            c["experience_years"],
            c["education_level"],
            c["recommendation"],
            c["status"],
            ", ".join(c["matched_skills"]),
            ", ".join(c["missing_skills"]),
            c["strengths"],
            c["weaknesses"]
        ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 40)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=candidate_rankings.xlsx"}
    )

@app.post("/clear")
async def clear(request: Request):
    candidates_store.clear()
    return templates.TemplateResponse(request, "screen.html", {
        "jobs": jobs_store,
        "results": [],
        "job_keywords": [],
        "job_description": ""
    })